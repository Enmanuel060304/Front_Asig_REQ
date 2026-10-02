"""Exportación de la tabla de asignación a Excel (una hoja, todas las columnas)."""
import tempfile
from datetime import date, datetime, time
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import repo

_NATIVOS = (int, float, Decimal, bool, datetime, date, time)


def _valor(v):
    if v is None or isinstance(v, _NATIVOS):
        return v
    # Texto y tipos no soportados (uuid, bytes…): a str sin caracteres de control que Excel rechaza
    return ILLEGAL_CHARACTERS_RE.sub("", str(v))


def asignacion_xlsx() -> tuple[str, int]:
    """Escribe la asignación en un .xlsx temporal. Devuelve (ruta, filas)."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Asignación")
    ws.freeze_panes = "A2"
    ncols = 0

    def columnas(nombres):
        nonlocal ncols
        ncols = len(nombres)
        encabezado = []
        for n in nombres:
            c = WriteOnlyCell(ws, value=n)
            c.font = Font(bold=True)
            encabezado.append(c)
        ws.append(encabezado)

    def filas(bloque):
        for f in bloque:
            ws.append([_valor(v) for v in f])

    total = repo.asignacion_exportar(columnas, filas)
    if ncols:
        ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{total + 1}"

    archivo = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    archivo.close()
    wb.save(archivo.name)
    return archivo.name, total
